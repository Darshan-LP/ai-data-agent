
"""
Export a pandas DataFrame to a formatted Excel workbook.
"""

from __future__ import annotations

import re
import uuid
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import get_settings


_SLUG_RE = re.compile(r"[^A-Za-z0-9_.-]+")


def _slug(text: str, fallback: str = "report") -> str:
    """Convert a title into a safe filename component."""
    text = (text or "").strip()
    text = _SLUG_RE.sub("_", text)[:60].strip("_.-")
    return text or fallback


def _outputs_dir() -> Path:
    """Create and return the configured output directory."""
    path = Path(get_settings().outputs_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def export_excel(
    df: pd.DataFrame,
    title: str = "Data Report",
) -> str:
    """
    Export a DataFrame to Excel and return the generated file path.

    Features:
    - Bold, colored column headers.
    - Frozen header row.
    - Automatic column widths.
    - Filters on the header row.
    """

    if df is None:
        raise ValueError("Cannot export a missing DataFrame.")

    if len(df.columns) == 0:
        raise ValueError("Cannot export a DataFrame without columns.")

    filename = f"{_slug(title)}-{uuid.uuid4().hex[:8]}.xlsx"
    output_path = _outputs_dir() / filename

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            index=False,
            sheet_name="Results",
        )

        worksheet = writer.sheets["Results"]

        # Freeze the first row so headers stay visible while scrolling.
        worksheet.freeze_panes = "A2"

        # Enable column filters.
        worksheet.auto_filter.ref = worksheet.dimensions

        # Style the header row.
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # Calculate readable column widths.
        for column_index, column_name in enumerate(
            df.columns,
            start=1,
        ):
            values = df[column_name].head(500).tolist()

            max_length = max(
                [len(str(column_name))]
                + [
                    len(str(value))
                    for value in values
                    if value is not None
                ]
            )

            width = min(max(max_length + 2, 12), 40)

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = width

        worksheet.row_dimensions[1].height = 22

    return str(output_path)
